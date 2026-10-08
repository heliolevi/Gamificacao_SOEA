import os
import re
import json
import base64
import io
import hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import FastAPI, HTTPException, Form, Query, UploadFile, File, Request, Depends
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import qrcode
from PIL import Image, ImageDraw, ImageFont
from datetime import datetime, date
import config
from database import banco_dados
from auth import get_current_user, require_admin
from auth_routes import router as auth_router, RegistroIn, CodigoEnviadoOut, registro_iniciar
from cache import cache_get, cache_set, valores_enum, RANKING_CACHE_TTL
from nivel import calcular_nivel
from vinculos import rotulo_vinculo

# uvicorn main:app --reload

app = FastAPI(
    title=f"{config.EVENTO_NOME} · Caça QR",
    # Documentação interativa só fora de produção (não expõe o mapa da API no evento)
    docs_url=None if config.EM_PRODUCAO else "/docs",
    redoc_url=None,
    openapi_url=None if config.EM_PRODUCAO else "/openapi.json",
)

FRONTEND_ORIGINS = [
    origem.strip()
    for origem in os.getenv(
        "FRONTEND_ORIGINS",
        f"http://localhost:3000,{config.PUBLIC_APP_URL}",
    ).split(",")
    if origem.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type", config.CSRF_HEADER],
)

app.include_router(auth_router)


@app.middleware("http")
async def cabecalhos_seguranca(request: Request, call_next):
    resposta = await call_next(request)
    resposta.headers.setdefault("X-Content-Type-Options", "nosniff")
    resposta.headers.setdefault("Referrer-Policy", "no-referrer")
    resposta.headers.setdefault("X-Frame-Options", "DENY")
    if request.url.path.startswith("/auth"):
        resposta.headers["Cache-Control"] = "no-store"
    return resposta

ENUMS_PERMITIDOS = {"escola", "curso_interesse",  "status_academico"}
PONTOS_POR_AMIGO = 50

@app.get("/opcoes/{nome}")
def opcoes(nome: str):
    """Busca os valores do enum diretamente do banco de dados (com cache curto)."""
    if nome not in ENUMS_PERMITIDOS:
        raise HTTPException(status_code=404, detail=f"Enum '{nome}' não encontrado.")

    return valores_enum(nome)

def validar_nome_sem_numeros(nome: str):
    if any(char.isdigit() for char in nome):
        return False, "O nome não pode conter números."
    if len(nome.strip()) < 6:
        return False, "Nome muito curto."
    return True, ""

def gerar_code_hash(semente: str) -> str:
    dados_hash = f"{semente}{os.urandom(8).hex()}"
    return hashlib.sha256(dados_hash.encode()).hexdigest()[:16]

@app.get("/")
def read_index():
    if config.EM_PRODUCAO:
        return {"status": "ok", "app": config.EVENTO_NOME}
    return RedirectResponse(url="/docs")

# ==================== USER ====================

@app.post("/usuarios/novo", response_model=CodigoEnviadoOut, status_code=202)
def cadastro_user(dados: RegistroIn, request: Request):
    """
    Mantida só por compatibilidade. NÃO cria conta: faz o mesmo que /auth/registro/iniciar
    (envia o código por e-mail). A conta só nasce em /auth/registro/verificar.
    """
    return registro_iniciar(dados, request)

@app.get("/usuarios/dados/exportar")
def exportar_dados(
    admin_id: str = Depends(require_admin),
    data: date = Query(...),
    formato: str = Query("xlsx"),
    pontos_min: int = Query(None),
    pontos_max: int = Query(None),
):
    query = (
        banco_dados.table("users")
        .select("nome", "pontos", "escola", 'curso_interesse', "status_academico", "vinculo")
        .gte("data_registro", f"{data}T00:00:00")
        .lte("data_registro", f"{data}T23:59:59")
        .eq("is_admin", False)
    )
    if pontos_min is not None:
        query = query.gte("pontos", pontos_min)
    if pontos_max is not None:
        query = query.lte("pontos", pontos_max)

    response = query.execute()
    if not response.data:
        raise HTTPException(status_code=404, detail=f"Nenhum usuário encontrado para os filtros informados.")

    usuarios = response.data

    if formato == "json":
        relatorio = {
            "data_relatorio": str(data),
            "filtros": { "pontos_min": pontos_min, "pontos_max": pontos_max},
            "total_alunos": len(usuarios),
            "alunos": usuarios
        }
        json_bytes = json.dumps(relatorio, ensure_ascii=False, indent=2).encode("utf-8")
        return StreamingResponse(io.BytesIO(json_bytes), media_type="application/json",
            headers={"Content-Disposition": f"attachment; filename=relatorio_{data}.json"})

    elif formato == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Relatório {data}"
        colunas = ["Nome", "Pontos", "Vínculo"]
        campos  = ["nome", "pontos", "vinculo"]
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for col_idx, titulo in enumerate(colunas, start=1):
            cell = ws.cell(row=1, column=col_idx, value=titulo)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_idx, usuario in enumerate(usuarios, start=2):
            for col_idx, campo in enumerate(campos, start=1):
                valor = usuario.get(campo, "")
                if campo == "vinculo":
                    valor = rotulo_vinculo(valor)
                ws.cell(row=row_idx, column=col_idx, value=valor)
        for col in ws.columns:
            max_length = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = max_length + 4
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return StreamingResponse(buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=relatorio_{data}.xlsx"})

    raise HTTPException(status_code=400, detail="Formato inválido.")

@app.get("/ranking")
def ranking(limit: int = 10):
    """
    Retorna o ranking ordenado por pontos (com cache curto pra aguentar picos).
    - limit=10  → top 10 para a tela de ranking/home (padrão)
    - limit=0   → todos os usuários
    """
    chave = f"ranking:{limit}"
    cache_hit = cache_get(chave, RANKING_CACHE_TTL)
    if cache_hit is not None:
        return cache_hit

    query = (
        banco_dados.table("users")
        .select("id_user, nome, pontos, catch(count), data_registro")
        .eq("is_admin", False)
        .order("pontos", desc=True)
        .order("data_registro", desc=False)
    )
    if limit > 0:
        query = query.limit(limit)

    res = query.execute()

    ranking_formatado = []
    for user in res.data:
        ranking_formatado.append({
            "id": user["id_user"],
            "nome": user["nome"],
            "pontos": user["pontos"],
            "nivel": calcular_nivel(user["pontos"]),
            "qrs_capturados": user["catch"][0]["count"] if user.get("catch") else 0
        })

    cache_set(chave, ranking_formatado)
    return ranking_formatado


@app.get("/usuarios/me/posicao")
def posicao_usuario(id_user: str = Depends(get_current_user)):
    """
    Retorna a posição, pontos e QRs do usuário autenticado sem carregar o ranking inteiro.
    Conta quantos usuários não-admin têm pontos maiores que o alvo.
    """
    usuario = (
        banco_dados.table("users")
        .select("id_user, nome, pontos, data_registro")
        .eq("id_user", id_user)
        .single()
        .execute()
    )
    if not usuario.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")

    pontos = usuario.data["pontos"]
    acima_por_pontos = (
        banco_dados.table("users")
        .select("id_user", count="exact")
        .eq("is_admin", False)
        .gt("pontos", pontos)
        .execute()
    )

    acima_por_registro = (
        banco_dados.table("users")
        .select("id_user", count="exact")
        .eq("is_admin", False)
        .eq("pontos", pontos)
        .lt("data_registro", usuario.data["data_registro"])
        .execute()
    )

    posicao = (acima_por_pontos.count or 0) + (acima_por_registro.count or 0) + 1

    qrs_evento = (
        banco_dados.table("catch")
        .select("id_catch", count="exact")
        .eq("id_user", id_user)
        .execute()
    )

    qrs_amigo = (
        banco_dados.table("friend_scans")
        .select("id", count="exact")
        .eq("scanner_id", id_user)
        .execute()
    )

    return {
        "id": id_user,
        "nome": usuario.data["nome"],
        "pontos": pontos,
        "nivel": calcular_nivel(pontos),
        "posicao": posicao,
        "qrs_capturados": (qrs_evento.count or 0) + (qrs_amigo.count or 0)
    }

@app.post("/responder")
def responder_pergunta(
    id_pergunta: str = Form(...),
    resposta: str = Form(...),
    tempo_segundos: int = Form(...),
    id_user: str = Depends(get_current_user),
):
    pergunta = banco_dados.table("perguntas").select("*").eq("id_pergunta", id_pergunta).single().execute()
    if not pergunta.data:
        raise HTTPException(status_code=404, detail="Pergunta não encontrada.")

    qrs_com_pergunta = (
        banco_dados.table("qrcodes")
        .select("code_hash")
        .eq("id_pergunta", id_pergunta)
        .execute()
    )
    hashes_validos = [q["code_hash"] for q in (qrs_com_pergunta.data or [])]

    if not hashes_validos:
        raise HTTPException(status_code=400, detail="Nenhum QR Code vinculado a esta pergunta.")

    captura_valida = (
        banco_dados.table("catch")
        .select("id_catch")
        .eq("id_user", id_user)
        .in_("code_hash", hashes_validos)
        .execute()
    )
    if not captura_valida.data:
        raise HTTPException(status_code=403, detail="Capture o QR Code antes de responder.")

    ja_respondeu = (
        banco_dados.table("user_perguntas")
        .select("id")
        .eq("id_user", id_user)
        .eq("id_pergunta", id_pergunta)
        .execute()
    )
    if ja_respondeu.data:
        raise HTTPException(status_code=400, detail="Você já respondeu esta pergunta.")

    p = pergunta.data
    acertou = resposta.upper() == p["resposta_correta"].upper()
    pontos_bonus = (p["pontos_rapido"] if tempo_segundos <= 10 else p["pontos_lento"]) if acertou else 0

    try:
        # Sempre registra, independente de acerto — o trigger cuida dos pontos
        banco_dados.table("user_perguntas").insert({
            "id_user": id_user,
            "id_pergunta": id_pergunta,
            "resposta": resposta,
            "tempo_segundos": tempo_segundos,
        }).execute()
    except Exception:
        raise HTTPException(status_code=500, detail="Erro ao registrar resposta. Tente novamente.")

    return {
        "acertou": acertou,
        "resposta_correta": p["resposta_correta"],
        "pontos_bonus": pontos_bonus,
        "feedback": "Resposta correta! 🎉" if acertou else "Resposta errada. Sem pontos bônus desta vez.",
    }

@app.patch("/usuarios/me/nome")
def atualizar_nome(nome: str = Form(...), id_user: str = Depends(get_current_user)):
    v_nome, m_nome = validar_nome_sem_numeros(nome)
    if not v_nome:
        raise HTTPException(status_code=400, detail=m_nome)

    resultado = banco_dados.table("users").update({"nome": nome.title()}).eq("id_user", id_user).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")
    return {"status": "Sucesso", "nome": resultado.data[0]["nome"]}

# ==================== QR PESSOAL / AMIGOS (XP) ====================

@app.get("/usuarios/me/qrcode")
def meu_qrcode(id_user: str = Depends(get_current_user)):
    usuario = banco_dados.table("users").select("personal_code_hash").eq("id_user", id_user).single().execute()
    if not usuario.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")

    personal_code_hash = usuario.data.get("personal_code_hash")
    if not personal_code_hash:
        personal_code_hash = gerar_code_hash(id_user)
        banco_dados.table("users").update({"personal_code_hash": personal_code_hash}).eq("id_user", id_user).execute()

    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(f"{config.PUBLIC_APP_URL}/perfil/{personal_code_hash}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return StreamingResponse(buf, media_type="image/png")

@app.post("/amigos/escanear")
def escanear_amigo(code_hash: str = Form(...), id_user: str = Depends(get_current_user)):
    alvo = banco_dados.table("users").select("id_user, nome").eq("personal_code_hash", code_hash).execute()
    if not alvo.data:
        raise HTTPException(status_code=404, detail="QR Code de amigo não encontrado.")

    scanned = alvo.data[0]
    scanned_id = scanned["id_user"]
    if scanned_id == id_user:
        raise HTTPException(status_code=400, detail="Você não pode escanear seu próprio QR Code.")

    # Networking é mútuo: bloqueia se essa dupla já se conectou em qualquer direção
    ja_conectados = (
        banco_dados.table("friend_scans")
        .select("id")
        .or_(
            f"and(scanner_id.eq.{id_user},scanned_id.eq.{scanned_id}),"
            f"and(scanner_id.eq.{scanned_id},scanned_id.eq.{id_user})"
        )
        .execute()
    )
    if ja_conectados.data:
        raise HTTPException(status_code=409, detail=f"Você e {scanned['nome']} já se conectaram antes.")

    try:
        banco_dados.table("friend_scans").insert({
            "scanner_id": id_user,
            "scanned_id": scanned_id,
        }).execute()
    except Exception as e:
        error_msg = str(e)
        if "23505" in error_msg or "duplicate key" in error_msg:
            raise HTTPException(status_code=409, detail=f"Você e {scanned['nome']} já se conectaram antes.")
        raise HTTPException(status_code=500, detail="Erro ao registrar a captura de amigo.")

    # Networking: os dois lados ganham pontos pela conexão
    scanner = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
    pontos_atual = scanner.data.get("pontos") or 0 if scanner.data else 0
    pontos_novo = pontos_atual + PONTOS_POR_AMIGO
    banco_dados.table("users").update({"pontos": pontos_novo}).eq("id_user", id_user).execute()

    alvo_atual = banco_dados.table("users").select("pontos").eq("id_user", scanned_id).single().execute()
    alvo_pontos_atual = alvo_atual.data.get("pontos") or 0 if alvo_atual.data else 0
    banco_dados.table("users").update({"pontos": alvo_pontos_atual + PONTOS_POR_AMIGO}).eq("id_user", scanned_id).execute()

    nivel_anterior = calcular_nivel(pontos_atual)
    nivel_atual = calcular_nivel(pontos_novo)

    return {
        "status": "Sucesso",
        "amigo": scanned["nome"],
        "pontos_ganho": PONTOS_POR_AMIGO,
        "pontos_total": pontos_novo,
        "nivel_anterior": nivel_anterior,
        "nivel_atual": nivel_atual,
        "subiu_de_nivel": nivel_atual > nivel_anterior,
    }

# ==================== QRCODE (EVENTO) ====================

@app.post("/qrcodes/gerar")
def gerar_qr(
    nome_local: str,
    pontos: int,
    admin_id: str = Depends(require_admin),
    id_pergunta: str = Query(None),
):
    code_hash = gerar_code_hash(f"{nome_local}-{pontos}")

    insert_data = {
        "code_hash": code_hash,
        "pontos": pontos,
        "local": nome_local,
    }
    if id_pergunta:
        insert_data["id_pergunta"] = id_pergunta

    banco_dados.table('qrcodes').insert(insert_data).execute()

    return {"status": "Sucesso", "code_hash": code_hash}


@app.get("/qrcodes/download/{code_hash}")
def download_qr(code_hash: str):
    response = banco_dados.table("qrcodes").select("local").eq("code_hash", code_hash).execute()
    if not response.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")

    local = response.data[0]["local"] or code_hash

    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(f"{config.PUBLIC_APP_URL}/scan/{code_hash}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    nome_arquivo = re.sub(r"[^a-zA-Z0-9_-]+", "_", local).strip("_") or code_hash
    return StreamingResponse(
        buf,
        media_type="image/png",
        headers={"Content-Disposition": f"attachment; filename=qrcode_{nome_arquivo}.png"},
    )

@app.get("/qrcodes/{code_hash}/pdf")
def download_qr_pdf(code_hash: str, admin_id: str = Depends(require_admin)):
    """Gera uma etiqueta pronta pra imprimir: QR Code + nome do local + pontuação."""
    response = banco_dados.table("qrcodes").select("local, pontos").eq("code_hash", code_hash).execute()
    if not response.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")

    local = response.data[0]["local"] or "QR Code"
    pontos = response.data[0]["pontos"]

    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(f"{config.PUBLIC_APP_URL}/scan/{code_hash}")
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")

    largura, altura = 600, 800
    etiqueta = Image.new("RGB", (largura, altura), "white")
    desenho = ImageDraw.Draw(etiqueta)

    fonte_titulo = ImageFont.load_default(size=32)
    fonte_legenda = ImageFont.load_default(size=24)

    def centralizar(texto, fonte, y):
        caixa = desenho.textbbox((0, 0), texto, font=fonte)
        x = (largura - (caixa[2] - caixa[0])) / 2
        desenho.text((x, y), texto, fill="black", font=fonte)

    centralizar(local, fonte_titulo, 40)

    qr_tamanho = 440
    qr_redimensionado = qr_img.resize((qr_tamanho, qr_tamanho))
    etiqueta.paste(qr_redimensionado, ((largura - qr_tamanho) // 2, 120))

    centralizar(f"+{pontos} pontos", fonte_legenda, 120 + qr_tamanho + 30)
    centralizar(f"{config.EVENTO_NOME} · Caça QR", fonte_legenda, altura - 60)

    buf = io.BytesIO()
    etiqueta.save(buf, format="PDF")
    buf.seek(0)

    nome_arquivo = re.sub(r"[^a-zA-Z0-9_-]+", "_", local).strip("_") or code_hash
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=qrcode_{nome_arquivo}.pdf"},
    )

@app.get("/qrcodes/listar")
def listar_qrcodes(admin_id: str = Depends(require_admin)):
    response = banco_dados.table('qrcodes').select("*").execute()
    return response.data

class StatusUpdate(BaseModel):
    ativo: bool

@app.patch("/qrcodes/status/{code_hash}")
def toggle_status_qr(code_hash: str, body: StatusUpdate, admin_id: str = Depends(require_admin)):
    resultado = banco_dados.table('qrcodes').update({"ativo": body.ativo}).eq("code_hash", code_hash).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")
    return {"status": "sucesso", "ativo": body.ativo}

@app.patch("/qrcodes/{code_hash}/vincular")
def vincular_qrcode(
    code_hash: str,
    admin_id: str = Depends(require_admin),
    id_pergunta: str = Query(None),
):
    atualizacao = {}
    if id_pergunta is not None:
        atualizacao["id_pergunta"] = None if id_pergunta == "null" else id_pergunta

    if not atualizacao:
        raise HTTPException(status_code=400, detail="Nenhum campo para atualizar.")

    resultado = banco_dados.table("qrcodes").update(atualizacao).eq("code_hash", code_hash).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado.")
    return {"status": "Sucesso", "qrcode": resultado.data[0]}


@app.post("/capturar")
def capturar(code_hash: str = Form(...), id_user: str = Depends(get_current_user)):

    try:

        qr_data = (
        banco_dados.table("qrcodes")
        .select("pontos, ativo, id_pergunta")
        .eq("code_hash", code_hash)
        .single()
        .execute()
        )
        if not qr_data.data:
            raise HTTPException(status_code=404, detail="QR Code não encontrado.")
        if not qr_data.data["ativo"]:
            raise HTTPException(status_code=403, detail="Este QR Code foi desativado.")
        # 1 Verifica se o QRcode existe ou se esta ativo


        ja_capturado = (
            banco_dados.table("catch")
            .select("id_catch")
            .eq("id_user", id_user)
            .eq("code_hash", code_hash)
            .execute()
        )
        # 2 verifica se ja foi capturado
        if ja_capturado.data:
            raise HTTPException(status_code=409, detail="Você já capturou este QR Code!")

        valor_pontos = qr_data.data["pontos"]
        id_pergunta  = qr_data.data["id_pergunta"]

        antes = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
        pontos_antes = (antes.data or {}).get("pontos") or 0

        # 3. Registra a captura do QR (o trigger trg_atualiza_pontos soma os pontos)
        banco_dados.table("catch").insert({
            "id_user": id_user,
            "catch_time": datetime.now().isoformat(),
            "code_hash": code_hash
        }).execute()

        depois = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
        pontos_depois = (depois.data or {}).get("pontos") or 0
        nivel_anterior = calcular_nivel(pontos_antes)
        nivel_atual = calcular_nivel(pontos_depois)

        # 6. Busca a pergunta vinculada, se houver
        pergunta = None
        if id_pergunta:
            p = banco_dados.table("perguntas").select("*").eq("id_pergunta", id_pergunta).single().execute()
            if p.data:
                pergunta = {
                    "id_pergunta": p.data["id_pergunta"],
                    "enunciado": p.data["enunciado"],
                    "tipo": p.data["tipo"],
                    "alternativas": p.data["alternativas"],
                }

        return {
            "status": "Sucesso",
            "pontos_qr": valor_pontos,
            "pontos_total": pontos_depois,
            "nivel_anterior": nivel_anterior,
            "nivel_atual": nivel_atual,
            "subiu_de_nivel": nivel_atual > nivel_anterior,
            "pergunta": pergunta,
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"Erro inesperado na captura: {e}")
        raise HTTPException(status_code=500, detail="Erro interno ao processar a captura. Tente novamente.")

# ==================== PERGUNTAS ====================

@app.post("/perguntas/nova")
def criar_pergunta(
    admin_id: str = Depends(require_admin),
    enunciado: str = Form(...),
    tipo: str = Form(...),
    resposta_correta: str = Form(...),
    pontos_rapido: int = Form(50),
    pontos_lento: int = Form(20),
    alternativa_a: str = Form(""),
    alternativa_b: str = Form(""),
    alternativa_c: str = Form(""),
    alternativa_d: str = Form(""),
):
    if tipo not in ("multipla_escolha", "verdadeiro_falso"):
        raise HTTPException(status_code=400, detail="Tipo inválido.")

    if tipo == "multipla_escolha":
        if not all([alternativa_a, alternativa_b, alternativa_c, alternativa_d]):
            raise HTTPException(status_code=400, detail="Preencha todas as alternativas A, B, C e D.")
        if resposta_correta not in ("A", "B", "C", "D"):
            raise HTTPException(status_code=400, detail="Resposta correta deve ser A, B, C ou D.")
        alternativas = {"A": alternativa_a, "B": alternativa_b, "C": alternativa_c, "D": alternativa_d}
    else:
        if resposta_correta not in ("V", "F"):
            raise HTTPException(status_code=400, detail="Resposta correta deve ser V ou F.")
        alternativas = {"V": "Verdadeiro", "F": "Falso"}

    resultado = banco_dados.table("perguntas").insert({
        "enunciado": enunciado,
        "tipo": tipo,
        "alternativas": alternativas,
        "resposta_correta": resposta_correta,
        "pontos_rapido": pontos_rapido,
        "pontos_lento": pontos_lento,
    }).execute()

    return {"status": "Sucesso", "pergunta": resultado.data[0]}

@app.get("/perguntas/listar")
def listar_perguntas(admin_id: str = Depends(require_admin)):
    response = banco_dados.table("perguntas").select("*").order("criado_em", desc=True).execute()
    return response.data or []

@app.patch("/perguntas/{id_pergunta}")
def editar_pergunta(
    id_pergunta: str,
    admin_id: str = Depends(require_admin),
    enunciado: str = Form(None),
    tipo: str = Form(None),
    resposta_correta: str = Form(None),
    pontos_rapido: int = Form(None),
    pontos_lento: int = Form(None),
    alternativa_a: str = Form(None),
    alternativa_b: str = Form(None),
    alternativa_c: str = Form(None),
    alternativa_d: str = Form(None),
):
    atualizacao = {}
    if enunciado is not None:
        atualizacao["enunciado"] = enunciado
    if tipo is not None:
        atualizacao["tipo"] = tipo
    if resposta_correta is not None:
        atualizacao["resposta_correta"] = resposta_correta
    if pontos_rapido is not None:
        atualizacao["pontos_rapido"] = pontos_rapido
    if pontos_lento is not None:
        atualizacao["pontos_lento"] = pontos_lento

    # Determina o tipo final (pode ter mudado nesta requisição ou já existia)
    tipo_final = tipo
    if tipo_final is None:
        current = banco_dados.table("perguntas").select("tipo, alternativas").eq("id_pergunta", id_pergunta).single().execute()
        tipo_final = current.data.get("tipo") if current.data else None

    if tipo_final == "verdadeiro_falso":
        # Sempre sobrescreve com apenas V/F, descartando qualquer A/B/C/D anterior
        atualizacao["alternativas"] = {"V": "Verdadeiro", "F": "Falso"}
    elif tipo_final == "multipla_escolha":
        if any(x is not None for x in [alternativa_a, alternativa_b, alternativa_c, alternativa_d]):
            # Busca alternativas atuais SÓ se o tipo não mudou (para preservar as que não foram enviadas)
            if tipo is None:  # tipo não mudou, faz merge
                current = banco_dados.table("perguntas").select("alternativas").eq("id_pergunta", id_pergunta).single().execute()
                alts = current.data.get("alternativas", {}) if current.data else {}
                # Limpa chaves de verdadeiro_falso que possam ter ficado
                alts = {k: v for k, v in alts.items() if k in ("A", "B", "C", "D")}
            else:  # tipo mudou para multipla_escolha, começa do zero
                alts = {}
            if alternativa_a is not None: alts["A"] = alternativa_a
            if alternativa_b is not None: alts["B"] = alternativa_b
            if alternativa_c is not None: alts["C"] = alternativa_c
            if alternativa_d is not None: alts["D"] = alternativa_d
            atualizacao["alternativas"] = alts

    if not atualizacao:
        raise HTTPException(status_code=400, detail="Nenhum campo para atualizar.")

    resultado = banco_dados.table("perguntas").update(atualizacao).eq("id_pergunta", id_pergunta).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="Pergunta não encontrada.")
    return {"status": "Sucesso", "pergunta": resultado.data[0]}

@app.delete("/perguntas/{id_pergunta}")
def deletar_pergunta(id_pergunta: str, admin_id: str = Depends(require_admin)):
    banco_dados.table("qrcodes").update({"id_pergunta": None}).eq("id_pergunta", id_pergunta).execute()
    banco_dados.table("perguntas").delete().eq("id_pergunta", id_pergunta).execute()
    return {"status": "Sucesso", "mensagem": "Pergunta removida."}
